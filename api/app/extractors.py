"""Document content extractors for various file formats and web sources."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import NamedTuple
from xml.etree import ElementTree

import httpx
from bs4 import BeautifulSoup
from pptx import Presentation
from pypdf import PdfReader


class ExtractedDocument(NamedTuple):
    source: str
    text: str


def extract_text_from_pdf(content: bytes, filename: str = "document.pdf") -> ExtractedDocument:
    """Extract text from PDF file content."""
    reader = PdfReader(BytesIO(content))
    texts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            texts.append(page_text.strip())
    return ExtractedDocument(source=filename, text="\n\n".join(texts))


def extract_text_from_pptx(content: bytes, filename: str = "document.pptx") -> ExtractedDocument:
    """Extract text from PowerPoint file content."""
    prs = Presentation(BytesIO(content))
    texts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_texts.append(shape.text.strip())
        if slide_texts:
            texts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
    return ExtractedDocument(source=filename, text="\n\n".join(texts))


def extract_text_from_txt(content: bytes, filename: str = "document.txt") -> ExtractedDocument:
    """Extract text from plain text file content."""
    text = content.decode("utf-8", errors="ignore")
    return ExtractedDocument(source=filename, text=text.strip())


def extract_text_from_md(content: bytes, filename: str = "document.md") -> ExtractedDocument:
    """Extract text from Markdown file content."""
    text = content.decode("utf-8", errors="ignore")
    return ExtractedDocument(source=filename, text=text.strip())


def extract_text_from_file(content: bytes, filename: str) -> ExtractedDocument:
    """Extract text from file based on extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(content, filename)
    elif ext in (".pptx", ".ppt"):
        return extract_text_from_pptx(content, filename)
    elif ext == ".md":
        return extract_text_from_md(content, filename)
    elif ext == ".txt":
        return extract_text_from_txt(content, filename)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


async def scrape_url(url: str, timeout: float = 30.0) -> ExtractedDocument:
    """Scrape text content from a URL."""
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
        headers = {
            "User-Agent": "Mozilla/5.0 (compatible; vo-chat/1.0; +https://github.com/vo-chat)"
        }
        response = await client.get(url, headers=headers)
        response.raise_for_status()

    content_type = response.headers.get("content-type", "").lower()
    if "application/pdf" in content_type:
        return extract_text_from_pdf(response.content, url)

    soup = BeautifulSoup(response.text, "lxml")

    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()

    main_content = soup.find("main") or soup.find("article") or soup.find("body")
    if main_content is None:
        main_content = soup

    text = main_content.get_text(separator="\n", strip=True)
    text = re.sub(r"\n{3,}", "\n\n", text)

    title = soup.find("title")
    title_text = title.get_text(strip=True) if title else url

    return ExtractedDocument(source=title_text or url, text=text)


async def scrape_sitemap_xml(xml_url: str, max_urls: int = 50, timeout: float = 30.0) -> list[ExtractedDocument]:
    """Parse sitemap XML and scrape URLs from it."""
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
        headers = {
            "User-Agent": "Mozilla/5.0 (compatible; vo-chat/1.0; +https://github.com/vo-chat)"
        }
        response = await client.get(xml_url, headers=headers)
        response.raise_for_status()

    root = ElementTree.fromstring(response.content)

    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    urls: list[str] = []

    for loc in root.findall(".//sm:loc", ns):
        if loc.text:
            urls.append(loc.text.strip())

    if not urls:
        for loc in root.findall(".//loc"):
            if loc.text:
                urls.append(loc.text.strip())

    urls = urls[:max_urls]

    documents: list[ExtractedDocument] = []
    for url in urls:
        try:
            doc = await scrape_url(url, timeout=timeout)
            if doc.text.strip():
                documents.append(doc)
        except Exception:
            continue

    return documents
